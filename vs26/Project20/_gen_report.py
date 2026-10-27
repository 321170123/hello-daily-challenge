# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

BASE = r'D:\新建文件夹\experiment5'

doc = Document()

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('实验五  汉字点阵的输出')
r.bold = True
r.font.size = Pt(16)

def head(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)

def body(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    r.font.size = Pt(12)

def pic(path, caption, width=Cm(14.5)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=width)
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.font.size = Pt(10.5)

head('一、Windows系统下用Visual Studio 2026实现')

body('老师给的printDot.c只能显示一个字，我把它改成能显示"许戈洋"三个字，字库文件还是原来的chs16.fon。')

pic(os.path.join(BASE, '许戈洋_VS2026.jpg'), '图1  Visual Studio 2026中修改后的源代码')

body('把chs16.fon放到工程目录下，按F5运行，控制台输出了"许戈洋"三个字的16x16点阵，'
     '1表示有点，0表示空白，能看出三个字的轮廓。')

pic(os.path.join(BASE, '许戈洋_Win.jpg'), '图2  Windows下程序运行输出的点阵结果')

head('二、Linux系统(Ubuntu16.04)下的实现')

body('通过VirtualBox的共享文件夹把代码和chs16.fon拷进Ubuntu虚拟机，在桌面新建experiment5文件夹保存。'
     'Linux下代码基本不变，只是去掉了system("pause")。')

pic(os.path.join(BASE, 'Linux_代码.jpg'), '图3  Ubuntu中用gedit打开的源代码')

body('在终端里用gcc编译后运行，输出了和Windows下一样的"许戈洋"三字点阵。')

pic(os.path.join(BASE, 'Linux_运行结果.jpg'), '图4  Linux终端中编译运行输出的点阵结果')

head('三、遇到的问题及解决方法')

body('1. 把.c文件直接拖进VS里按F5运行不了，弹出"附加到进程"对话框。'
     '原因是单个源文件不属于任何工程，要先打开解决方案再运行。')
body('2. 往虚拟机里拷文件时cp提示"没有那个文件或目录"，是共享文件夹里还没放实验文件，复制过去就好了。')
body('3. 编译时提示找不到printdot.c，但文件明明在。原因是Linux文件名区分大小写，'
     '实际文件名是printDot.c，改对大小写后编译通过。')
body('4. 原版程序里的system("pause")在Linux下提示not found，因为pause是Windows的命令，删掉即可。')

out = os.path.join(BASE, 'Test5_26020007114_许戈洋.docx')
doc.save(out)
print('saved:', out)
